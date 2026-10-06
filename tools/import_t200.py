"""Reproduce the bundled thruster profile from the unchanged official workbook.

Run: python tools/import_t200.py
Source S6. Normalization u=(PWM-1500)/400; kgf->N uses standard g=9.80665.
Conservative cumulative maximum envelopes remove small measurement inversions,
retaining all PWM points and no fitted arbitrary power exponent.
"""

import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "sources/T200-Public-Performance-Data-10-20V-September-2019.xlsx"
TARGET = ROOT / "holoenergy/profiles/thrusters/bluerobotics_t200.json"


def build():
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    tables = []
    changed = 0
    discrepancies = []
    for voltage in (10, 12, 14, 16, 18, 20):
        sheet = workbook[f"{voltage} V"]
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        table = {"voltage_V": voltage}
        for direction, sign in (("forward", 1), ("reverse", -1)):
            branch = sorted(
                (r for r in rows if (r[0] - 1500) * sign >= 0),
                key=lambda r: abs(r[0] - 1500),
            )
            force_envelope = power_envelope = 0.0
            data = []
            for row in branch:
                pwm, _, current, measured_v, power, force_kgf = row[:6]
                if abs(power - current * measured_v) > 1e-8:
                    discrepancies.append(
                        {
                            "sheet": sheet.title,
                            "PWM_us": pwm,
                            "published_power_W": power,
                            "current_times_voltage_W": current * measured_v,
                        }
                    )
                command = abs(pwm - 1500) / 400
                force = abs(force_kgf) * 9.80665
                if command == 0:
                    if abs(power) > 1e-8 or force > 1e-8:
                        raise ValueError("Neutral workbook point is not zero")
                new_force, new_power = max(force_envelope, force), max(power_envelope, power)
                changed += new_force != force or new_power != power
                force_envelope, power_envelope = new_force, new_power
                data.append(
                    {
                        "command": round(command, 8),
                        "force_N": round(force_envelope, 10),
                        "power_W": round(power_envelope, 10),
                    }
                )
            table[direction] = data
        tables.append(table)
    profile = {
        "model": "BlueRobotics_T200_2019",
        "metadata": {
            "source_id": "S6",
            "value_kind": "manufacturer_test_data_processed",
            "characterization": "static/bollard",
            "limitations": "The current T200 energy model is based on static/bollard manufacturer characterization and does not explicitly model inflow-dependent propeller performance, vehicle-speed effects, thruster-thruster interaction or installation effects.",
            "url": "https://cad.bluerobotics.com/T200-Public-Performance-Data-10-20V-September-2019.xlsx",
            "retrieved_date": "2026-10-06",
            "sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
            "processing": "PWM normalization; kgf to N; cumulative max force/power per direction",
            "points_adjusted_by_envelope": changed,
            "source_discrepancies": discrepancies,
            "discrepancy_policy": "Retain published Power column; flagged for manufacturer clarification",
            "calibrate": ["installed inflow/interference", "PWM mapping and ESC revision"],
            "electrical_boundary": "manufacturer measured electrical input; do not double-count ESC efficiency",
        },
        "voltage_tables": tables,
    }
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    TARGET.write_text(json.dumps(profile, indent=2) + "\n", encoding="utf-8")
    print(f"Generated {TARGET.name}: {len(tables)} voltages; {changed} envelope adjustments")


if __name__ == "__main__":
    build()
